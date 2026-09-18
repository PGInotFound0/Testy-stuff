import re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

src, dst = sys.argv[1], sys.argv[2]
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
for s in doc.sections:
    s.top_margin = s.bottom_margin = Cm(2)
    s.left_margin = s.right_margin = Cm(2.2)

def add_text(p, text):
    text = text.replace("`", "")
    for i, part in enumerate(text.split("**")):
        if not part:
            continue
        run = p.add_run(part)
        run.bold = bool(i % 2)
        run.italic = False
    return p

def add_para(text, style=None):
    p = doc.add_paragraph(style=style)
    return add_text(p, text)

for raw in open(src, encoding="utf-8").read().splitlines():
    line = raw.rstrip()
    if not line.strip():
        continue
    if line.startswith("# "):
        doc.add_heading(line[2:].strip(), level=0)
    elif line.startswith("## "):
        doc.add_heading(line[3:].strip(), level=1)
    elif line.startswith("### "):
        doc.add_heading(line[4:].strip(), level=2)
    elif line.strip() == "---":
        doc.add_paragraph("_" * 60)
    elif re.match(r"^\s*[-*] ", line):
        indent = len(line) - len(line.lstrip())
        p = doc.add_paragraph(style="List Bullet")
        add_text(p, line.strip()[2:].strip())
        p.paragraph_format.left_indent = Cm(0.6 + 0.6 * (indent // 2))
    elif re.match(r"^\s*\d+\.\s", line):
        add_para(line.strip(), style="List Number")
    else:
        add_para(line.strip())
doc.save(dst)
print("written", dst)
